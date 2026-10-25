"""
Gera o trajeto.json (dados do trajeto) a partir da planilha da Giulia.
Rascunho: saida-alvo provisoria = L-0-01 (Entrada do terreo).
Rotas = menor caminho ate a saida (calculado AQUI, so para gerar o rascunho;
o app nao calcula rota). Curvas vem da geometria das plantas (coordenadas em pixels).
Uso: python3 gerar_trajeto.py planilha.xlsx saida.json
"""
import sys, json, math, heapq
import openpyxl

EXIT = "L-0-01"
STAIRS_M = 13.2   # escada terreo<->1o andar (a confirmar: projecao horizontal)
RAMP_TOTAL = {"L-0": 6.1, "L-1": 6.3, "L-2": 6.3}   # E4 -> rampa -> P3 (divisao nao informada)

# coordenadas aproximadas lidas das plantas (so para calcular curvas)
XY = {
 # terreo
 "L-0-01": (158,591), "L-0-02": (125,521), "L-0-03": (391,484), "L-0-04": (524,484),
 "L-0-05": (675,484), "L-0-06": (824,484), "L-0-07": (1033,484), "L-0-09": (1033,217),
 "L-0-10": (858,249), "L-0-11": (708,249), "L-0-12": (566,249), "L-0-13": (417,249),
 "L-0-14": (274,249), "L-0-15": (274,391),
 # 1o andar
 "L-1-01": (274,441), "L-1-02": (125,521), "L-1-03": (391,484), "L-1-04": (524,484),
 "L-1-05": (675,484), "L-1-06": (824,484), "L-1-07": (1033,484), "L-1-09": (1033,217),
 "L-1-10": (858,249), "L-1-11": (708,249), "L-1-12": (566,249), "L-1-13": (417,249),
 "L-1-14": (274,249),
 # 2o andar
 "L-2-01": (274,441), "L-2-02": (125,521), "L-2-03": (391,484), "L-2-04": (524,484),
 "L-2-05": (675,484), "L-2-06": (824,484), "L-2-07": (1033,484), "L-2-09": (1033,217),
 "L-2-10": (708,249), "L-2-11": (566,249), "L-2-12": (417,249), "L-2-13": (104,249),
}

def heading(a, b):
    (x1,y1),(x2,y2) = XY[a], XY[b]
    return math.degrees(math.atan2(x2-x1, -(y2-y1)))   # 0 = cima, horario positivo

def turn(prev, node, nxt, thr=25):
    d = (heading(node, nxt) - heading(prev, node) + 540) % 360 - 180
    if d > thr: return "right"
    if d < -thr: return "left"
    return "none"

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
pts, edges = {}, {}

def add_edge(a, b, m):
    edges.setdefault(a, {})[b] = m
    edges.setdefault(b, {})[a] = m

def read(floor_prefix, sp, sc):
    for r in wb[sp].iter_rows(min_row=3, values_only=True):
        if r[1] and str(r[1]).startswith("L-") and r[3] and "Retorno" not in str(r[3]):
            pts[r[1]] = dict(tipo=str(r[3]).strip(), desc=r[4])
    ramp = RAMP_TOTAL[floor_prefix]
    for r in wb[sc].iter_rows(min_row=4, values_only=True):
        a, b, m = r[1], r[2], r[3]
        if not a or not b or not str(a).startswith(floor_prefix): continue
        if not isinstance(m, (int, float)): continue   # trechos da rampa (sem divisao)
        add_edge(a, b, float(m))
    n = lambda k: f"{floor_prefix}-{k:02d}"
    add_edge(n(7), n(9), ramp)        # E4 -> P3 (rampa fundida)
    pts.pop(n(8), None)               # ponto da rampa sem QR

read("L-0", "Térreo - Pontos", "Térreo - Conexões")
read("L-1", "Pontos", "Conexões")
read("L-2", "2º andar - Pontos", "2º andar - Conexões")
add_edge("L-1-01", "L-0-15", STAIRS_M)
add_edge("L-2-01", "L-1-01", STAIRS_M)

def shortest(src):
    dist, prev, pq = {src: 0}, {}, [(0, src)]
    while pq:
        d, u = heapq.heappop(pq)
        if d > dist[u]: continue
        for v, m in edges[u].items():
            if d + m < dist.get(v, 1e9):
                dist[v], prev[v] = d + m, u
                heapq.heappush(pq, (d + m, v))
    return prev

routes = []
for start in sorted(pts):
    if start == EXIT: continue
    prev = shortest(EXIT)             # caminhos a partir da saida (grafo nao direcionado)
    path, u = [start], start
    while u != EXIT:
        u = prev[u]; path.append(u)
    legs = []
    for i in range(len(path) - 1):
        a, b = path[i], path[i + 1]
        stairs = {a, b} in ({"L-1-01", "L-0-15"}, {"L-2-01", "L-1-01"})
        t = "none" if (i == len(path) - 2 or stairs) else turn(a, b, path[i + 2])
        leg = {"toId": b, "meters": edges[a][b], "turn": t}
        if stairs: leg["kind"] = "stairs"
        legs.append(leg)
    routes.append({"startId": start, "legs": legs})

points = []
for pid in sorted(pts, key=lambda k: (k[2], k)):
    p = pts[pid]
    points.append({"id": pid, "description": p["desc"], "noElevator": True,
                   "extinguisher": "neste ponto" if p["tipo"] == "Extintor" else ""})
out = {"exitId": EXIT, "exitName": "Entrada do térreo", "points": points, "routes": routes}
json.dump(out, open(sys.argv[2], "w", encoding="utf-8"), ensure_ascii=False, indent=2)
print(len(points), "pontos,", len(routes), "rotas")
